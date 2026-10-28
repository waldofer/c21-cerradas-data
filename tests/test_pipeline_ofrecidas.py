import sys

import pytest

import actualizar_cerradas as ac


def test_url_vendidas_no_cambia():
    assert ac.construir_url_busqueda("2026-10-01", "2026-10-19") == (
        "https://21online.century21.com.py/inventarios?oficinaStatus=todas&oficina="
        "&estatus=cerrada&tipoOperacion=venta&mesCierreDesde=2026-10-01"
        "&mesCierreHasta=2026-10-19&cantidad=30000&conFoto=1"
    )


def test_url_ofrecidas():
    url = ac.construir_url_busqueda("2015-01-01", "2026-10-19", estatus="enPromocion")
    assert "estatus=enPromocion" in url
    assert "fechaAltaDesde=2015-01-01" in url
    assert "fechaAltaHasta=2026-10-19" in url
    assert "conFoto=1" in url
    assert "tipoOperacion=venta" in url
    assert "mesCierre" not in url


def _celdas(ruta, hoja):
    import openpyxl

    return [c for f in openpyxl.load_workbook(ruta)[hoja].iter_rows(values_only=True) for c in f]


def test_main_corre_ofrecidas_aunque_vendidas_aborte(monkeypatch, tmp_path):
    resumen = tmp_path / "resumen.xlsx"
    dataset = tmp_path / "cerradas.json"
    dataset.write_text('[{"id": 1}, {"id": 2}, {"id": 3}]', encoding="utf-8")
    llamadas = []

    def vendidas_rota():
        raise Exception("fallo login")

    def ofrecidas():
        # El Resumen con la alerta ya tiene que estar escrito cuando corren las ofrecidas.
        llamadas.append(("ofrecidas", resumen.exists()))

    monkeypatch.setattr(sys, "argv", ["actualizar_cerradas.py"])
    monkeypatch.setattr(ac, "RESUMEN_PATH", resumen)
    monkeypatch.setattr(ac, "DATASET_PATH", dataset)
    monkeypatch.setattr(ac, "actualizar_vendidas", vendidas_rota)
    monkeypatch.setattr(ac, "actualizar_ofrecidas", ofrecidas)
    with pytest.raises(SystemExit) as salida:
        ac.main()
    assert salida.value.code == 1
    assert llamadas == [("ofrecidas", True)]
    celdas = _celdas(resumen, "Resumen")
    assert any(isinstance(c, str) and "Fallo vendidas: Exception: fallo login" in c for c in celdas)
    assert 3 in celdas  # total acumulado del cerradas.json actual


def test_main_alerta_vendidas_sin_valores_de_celdas(monkeypatch, tmp_path):
    resumen = tmp_path / "resumen.xlsx"

    def vendidas_rota():
        raise ValueError("invalid literal for int() with base 10: 'Juan Perez 0981'")

    monkeypatch.setattr(sys, "argv", ["actualizar_cerradas.py"])
    monkeypatch.setattr(ac, "RESUMEN_PATH", resumen)
    monkeypatch.setattr(ac, "DATASET_PATH", tmp_path / "no_existe.json")
    monkeypatch.setattr(ac, "actualizar_vendidas", vendidas_rota)
    monkeypatch.setattr(ac, "actualizar_ofrecidas", lambda: None)
    with pytest.raises(SystemExit):
        ac.main()
    celdas = _celdas(resumen, "Resumen")
    alerta = next(c for c in celdas if isinstance(c, str) and "Fallo vendidas" in c)
    assert "ValueError" in alerta
    assert "Juan" not in alerta and "0981" not in alerta
    assert 0 in celdas  # sin cerradas.json, total 0


def test_main_falla_ofrecidas_sale_con_1(monkeypatch):
    monkeypatch.setattr(sys, "argv", ["actualizar_cerradas.py", "--solo-ofrecidas"])
    monkeypatch.setattr(ac, "actualizar_ofrecidas", lambda: False)
    with pytest.raises(SystemExit) as salida:
        ac.main()
    assert salida.value.code == 1


def test_main_todo_ok_no_sale_con_error(monkeypatch):
    llamadas = []
    monkeypatch.setattr(sys, "argv", ["actualizar_cerradas.py"])
    monkeypatch.setattr(ac, "actualizar_vendidas", lambda: llamadas.append("vendidas"))
    monkeypatch.setattr(ac, "actualizar_ofrecidas", lambda: llamadas.append("ofrecidas") or True)
    ac.main()  # no levanta SystemExit
    assert llamadas == ["vendidas", "ofrecidas"]


def test_main_solo_ofrecidas(monkeypatch):
    llamadas = []
    monkeypatch.setattr(sys, "argv", ["actualizar_cerradas.py", "--solo-ofrecidas"])
    monkeypatch.setattr(ac, "actualizar_vendidas", lambda: llamadas.append("vendidas"))
    monkeypatch.setattr(ac, "actualizar_ofrecidas", lambda: llamadas.append("ofrecidas"))
    ac.main()
    assert llamadas == ["ofrecidas"]


def test_vendidas_publicadas_lee_head(monkeypatch):
    comandos = []

    class R:
        stdout = '[{"id": 7, "pr": 90000.0, "dir": "Peña 1"}]'.encode("utf-8")

    def falso_run(cmd, **kw):
        comandos.append((cmd, kw))
        return R()

    monkeypatch.setattr(ac.subprocess, "run", falso_run)
    assert ac.vendidas_publicadas() == {7: {"id": 7, "pr": 90000.0, "dir": "Peña 1"}}
    cmd, kw = comandos[0]
    assert cmd == ["git", "show", "HEAD:data/cerradas.json"]
    assert kw["cwd"] == ac.REPO_DIR and kw["check"] is True and kw["capture_output"] is True


def test_git_ofrecidas_stagea_solo_su_archivo(monkeypatch):
    comandos = []

    class R:
        returncode = 1  # hay diff

    monkeypatch.setattr(ac.subprocess, "run", lambda cmd, **kw: comandos.append(cmd) or R())
    ac.git_commit_push_ofrecidas("msg")
    assert ["git", "add", "data/ofrecidas.json"] in comandos
    assert ["git", "push"] in comandos
    for cmd in comandos:
        assert "-A" not in cmd and "." not in cmd and "--all" not in cmd


def test_git_ofrecidas_sin_diff_no_commitea(monkeypatch):
    comandos = []

    class R:
        returncode = 0

    monkeypatch.setattr(ac.subprocess, "run", lambda cmd, **kw: comandos.append(cmd) or R())
    ac.git_commit_push_ofrecidas("msg")
    assert not any(c[:2] in (["git", "commit"], ["git", "push"]) for c in comandos)


def test_actualizar_ofrecidas_excepcion_queda_como_alerta(monkeypatch, tmp_path):
    import openpyxl

    resumen = tmp_path / "resumen.xlsx"
    tmp_export = tmp_path / "export.xlsx"
    tmp_export.write_bytes(b"x")

    def sin_navegador():
        raise RuntimeError("sin red")

    monkeypatch.setattr(ac, "RESUMEN_PATH", resumen)
    monkeypatch.setattr(ac, "DESCARGA_OFRECIDAS_TMP", tmp_export)
    monkeypatch.setattr(ac, "sync_playwright", sin_navegador)
    assert ac.actualizar_ofrecidas() is False  # no debe levantar; avisa que fallo
    celdas = [c for f in openpyxl.load_workbook(resumen)["Ofrecidas"].iter_rows(values_only=True) for c in f]
    assert any(isinstance(c, str) and c.startswith("ALERTA - ") and "sin red" in c for c in celdas)
    assert not tmp_export.exists()


def test_ofrecidas_publicadas_lee_head(monkeypatch):
    comandos = []

    class R:
        def __init__(self, stdout):
            self.stdout = stdout.encode("utf-8")

    def falso_run(cmd, **kw):
        comandos.append((cmd, kw))
        if cmd[:2] == ["git", "ls-tree"]:
            return R("data/ofrecidas.json\n")
        return R('[{"id": 3, "dir": "Peña 2"}]')

    monkeypatch.setattr(ac.subprocess, "run", falso_run)
    assert ac.ofrecidas_publicadas() == [{"id": 3, "dir": "Peña 2"}]
    cmd, kw = comandos[-1]
    assert cmd == ["git", "show", "HEAD:data/ofrecidas.json"]
    assert kw["cwd"] == ac.REPO_DIR and kw["check"] is True and kw["capture_output"] is True


def test_ofrecidas_publicadas_sin_archivo_en_head(monkeypatch):
    comandos = []

    class R:
        stdout = b""

    monkeypatch.setattr(ac.subprocess, "run", lambda cmd, **kw: comandos.append(cmd) or R())
    assert ac.ofrecidas_publicadas() == []
    assert not any(c[:2] == ["git", "show"] for c in comandos)


class _NavegadorFalso:
    def __enter__(self):
        return self

    def __exit__(self, *a):
        return False

    @property
    def chromium(self):
        return self

    def launch(self, **kw):
        return self

    def new_context(self):
        return self

    def new_page(self):
        return self

    def close(self):
        pass


def test_actualizar_ofrecidas_con_alerta_no_escribe_ni_pushea(monkeypatch, tmp_path):
    resumen = tmp_path / "resumen.xlsx"
    tmp_export = tmp_path / "export.xlsx"
    llamadas = []

    def exportar(context, hoy, destino=None):
        tmp_export.write_bytes(b"x")
        return tmp_export

    monkeypatch.setattr(ac, "RESUMEN_PATH", resumen)
    monkeypatch.setattr(ac, "DESCARGA_OFRECIDAS_TMP", tmp_export)
    monkeypatch.setattr(ac, "sync_playwright", _NavegadorFalso)
    monkeypatch.setattr(ac, "login", lambda page: None)
    monkeypatch.setattr(ac, "exportar_ofrecidas_por_tramos", exportar)
    monkeypatch.setattr(ac, "ofrecidas_publicadas", lambda: [])
    monkeypatch.setattr(ac, "vendidas_publicadas", lambda: {})
    monkeypatch.setattr(ac, "procesar_ofrecidas", lambda *a: (None, None, "descarga rara"))
    monkeypatch.setattr(ac, "guardar_ofrecidas", lambda *a, **k: llamadas.append("guardar"))
    monkeypatch.setattr(ac, "git_commit_push_ofrecidas", lambda *a, **k: llamadas.append("push"))
    assert ac.actualizar_ofrecidas() is True  # la guarda funciono: no es una falla de la corrida
    assert llamadas == []
    assert "ALERTA - descarga rara" in _celdas(resumen, "Ofrecidas")
    assert not tmp_export.exists()


# --- Descarga de ofrecidas por tramos de fecha de alta ---
# La corrida del 2026-10-26 se cayo ("Page crashed") pidiendo las ~8.400 de una vez.

from datetime import date, timedelta

import openpyxl
import pytest


def test_tramos_ofrecidas_cubren_todo_sin_huecos():
    tramos = ac.tramos_ofrecidas(date(2026, 10, 28))
    assert tramos[0][0] == "2015-01-01"
    assert tramos[-1][1] == "2026-10-28"
    for (_, hasta), (desde, _) in zip(tramos, tramos[1:]):
        assert date.fromisoformat(desde) == date.fromisoformat(hasta) + timedelta(days=1)
    desde_2025 = [t for t in tramos if t[0] >= "2025-01-01"]
    assert all(date.fromisoformat(h) - date.fromisoformat(d) <= timedelta(days=62) for d, h in desde_2025)


def test_tramos_ofrecidas_mes_en_curso_va_con_el_anterior():
    # a principio de mes no se pide un tramo de un par de dias que podria venir vacio
    assert ac.tramos_ofrecidas(date(2026, 11, 2))[-1] == ("2026-10-01", "2026-11-02")
    assert ac.tramos_ofrecidas(date(2025, 1, 5))[-1] == ("2025-01-01", "2025-01-05")


class _ContextoFalso:
    def __init__(self):
        self.abiertas = self.cerradas = 0

    def new_page(self):
        self.abiertas += 1
        ctx = self

        class _Pagina:
            def close(self):
                ctx.cerradas += 1
        return _Pagina()


def _exportador(fallas=(), vacio=(), header=("ID", "PRECIO")):
    """Escribe un xlsx por tramo con una fila cuyo ID sale del tramo. fallas: {desde: cuantas veces falla}."""
    pendientes = dict(fallas)
    pedidos = []

    def exportar(page, desde, hasta, estatus, destino, **kw):
        pedidos.append((desde, hasta))
        assert estatus == "enPromocion"
        if pendientes.get(desde):
            pendientes[desde] -= 1
            raise RuntimeError("Page.goto: Page crashed")
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(list(header) if desde != "2024-01-01" or header == ("ID", "PRECIO") else ["ID", "OTRA"])
        if desde not in vacio:
            ws.append([desde.replace("-", ""), "1,00 USD"])
        wb.save(destino)
        return destino
    return exportar, pedidos


def _filas(path):
    return list(openpyxl.load_workbook(path, read_only=True).active.iter_rows(values_only=True))


def test_por_tramos_junta_todo_y_reintenta(monkeypatch, tmp_path):
    exportar, pedidos = _exportador(fallas={"2025-03-01": 1})
    monkeypatch.setattr(ac, "exportar_reporte", exportar)
    monkeypatch.setattr(ac, "DESCARGA_OFRECIDAS_TRAMO_TMP", tmp_path / "tramo.xlsx")
    ctx = _ContextoFalso()
    destino = ac.exportar_ofrecidas_por_tramos(ctx, date(2026, 10, 28), destino=tmp_path / "todo.xlsx")
    tramos = ac.tramos_ofrecidas(date(2026, 10, 28))
    filas = _filas(destino)
    assert filas[0] == ("ID", "PRECIO")
    assert len(filas) - 1 == len(tramos)  # una fila por tramo, ninguno perdido
    assert len(pedidos) == len(tramos) + 1  # el tramo que fallo se pidio dos veces
    assert ctx.abiertas == ctx.cerradas == len(pedidos)  # una pagina nueva por intento, siempre cerrada
    assert not (tmp_path / "tramo.xlsx").exists()


def test_por_tramos_si_un_tramo_no_baja_falla_todo(monkeypatch, tmp_path):
    exportar, _ = _exportador(fallas={"2025-03-01": 5})
    monkeypatch.setattr(ac, "exportar_reporte", exportar)
    monkeypatch.setattr(ac, "DESCARGA_OFRECIDAS_TRAMO_TMP", tmp_path / "tramo.xlsx")
    with pytest.raises(RuntimeError, match="2025-03-01"):
        ac.exportar_ofrecidas_por_tramos(_ContextoFalso(), date(2026, 10, 28), destino=tmp_path / "todo.xlsx")
    assert not (tmp_path / "todo.xlsx").exists()
    assert not (tmp_path / "tramo.xlsx").exists()


def test_por_tramos_un_tramo_vacio_es_falla(monkeypatch, tmp_path):
    exportar, _ = _exportador(vacio={"2025-06-01"})
    monkeypatch.setattr(ac, "exportar_reporte", exportar)
    monkeypatch.setattr(ac, "DESCARGA_OFRECIDAS_TRAMO_TMP", tmp_path / "tramo.xlsx")
    with pytest.raises(RuntimeError, match="2025-06-01"):
        ac.exportar_ofrecidas_por_tramos(_ContextoFalso(), date(2026, 10, 28), destino=tmp_path / "todo.xlsx")


def test_por_tramos_columnas_distintas_es_falla(monkeypatch, tmp_path):
    exportar, _ = _exportador(header=("ID", "PRECIO", "X"))
    monkeypatch.setattr(ac, "exportar_reporte", exportar)
    monkeypatch.setattr(ac, "DESCARGA_OFRECIDAS_TRAMO_TMP", tmp_path / "tramo.xlsx")
    with pytest.raises(RuntimeError, match="columnas"):
        ac.exportar_ofrecidas_por_tramos(_ContextoFalso(), date(2026, 10, 28), destino=tmp_path / "todo.xlsx")
