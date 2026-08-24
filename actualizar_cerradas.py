"""
Actualizador semanal del dataset de cerradas - Century 21 Paraguay
====================================================================
Corre los lunes (Task Scheduler), exporta lo cerrado en 21online desde la
última actualización hasta hoy, y actualiza (upsert por ID) el dataset
maestro en el repo de GitHub c21-cerradas-data.

REQUIERE ENGANCHARSE A TU LOGIN EXISTENTE:
Reemplazá la función `login()` por la misma que ya usás en infocasas-bot
(scrape.py) para entrar a 21online.century21.com.py. No la reescribas,
reusala tal cual — esto vive en un repo aparte, no toques infocasas-bot.

Flujo:
1. Lee data/cerradas.json del repo local (el dataset acumulado hasta ahora)
2. Calcula la fecha de corte: la fecha de cierre más reciente que ya tenemos,
   menos 14 días de margen (para agarrar cierres que se cargaron tarde)
3. Entra a 21online, exporta el reporte de cerradas filtrado por ESE rango
   (fecha_corte -> hoy)
4. Parsea el Excel descargado (misma lógica de limpieza que ya usamos:
   USD/PYG, LAT/LON, PRECIO vs PRECIO CIERRE, FECHA ALTA)
5. Upsert por ID sobre el dataset acumulado (los IDs nuevos se agregan,
   los existentes se actualizan si cambiaron)
6. Commit + push al repo c21-cerradas-data

Uso:
    python actualizar_cerradas.py

Requiere: playwright, openpyxl, git configurado con acceso de push al repo.
"""

import json
import re
import subprocess
import os
from datetime import datetime, timedelta
from pathlib import Path

import openpyxl
from playwright.sync_api import sync_playwright

from clasificar_precios import clasificar, construir_pool, UMBRAL_GS
from cotizaciones_bcp import TABLA_PATH, descargar_anho

REPO_DIR = Path(__file__).parent  # asumido: este script vive adentro del repo c21-cerradas-data
DATASET_PATH = REPO_DIR / "data" / "cerradas.json"
DESCARGA_TMP = REPO_DIR / "_tmp_export.xlsx"
RESUMEN_PATH = REPO_DIR / "resumen_actualizacion.xlsx"

TC_APROX = 6100  # Gs por USD, fallback si algún PRECIO viene en PYG sin más contexto
MARGEN_DIAS = 14  # relee los últimos N días por si hubo cargas tardías
LIMITE_NUEVAS = 500  # por encima de esto, algo salio mal (rango mal calculado, export duplicado, etc.)
COLUMNAS_REQUERIDAS = {"ID", "PRECIO CIERRE", "FECHA CIERRE", "M2T", "MUNICIPIO"}


BASE_URL = "https://21online.century21.com.py"


def construir_url_busqueda(fecha_desde: str, fecha_hasta: str) -> str:
    params = {
        "oficinaStatus": "todas",
        "oficina": "",
        "estatus": "cerrada",
        "tipoOperacion": "venta",
        "mesCierreDesde": fecha_desde,
        "mesCierreHasta": fecha_hasta,
        "cantidad": "30000",
        "conFoto": "1",
    }
    query = "&".join(f"{k}={v}" for k, v in params.items())
    return f"{BASE_URL}/inventarios?{query}"


def login(page, max_intentos: int = 3):
    """
    Login en 21online.century21.com.py. Adaptado de la función _login() de
    infocasas-bot/scrape.py (mismo sitio, mismo formulario): llena los dos
    primeros inputs visibles del form (usuario/contraseña) y clickea el
    submit, con reintentos si la sesión sigue en /login.

    Credenciales SIEMPRE desde variables de entorno, nunca hardcodeadas ni
    pegadas en el chat:
        C21_USUARIO, C21_PASSWORD
    """
    usuario = os.environ["C21_USUARIO"]
    password = os.environ["C21_PASSWORD"]

    for intento in range(1, max_intentos + 1):
        page.goto(f"{BASE_URL}/", wait_until="domcontentloaded", timeout=30000)
        page.wait_for_timeout(3000)

        inputs = page.locator("input:visible").all()
        if len(inputs) >= 1:
            inputs[0].fill(usuario)
        if len(inputs) >= 2:
            inputs[1].fill(password)
        else:
            try:
                page.locator("input[type='password']").first.fill(password)
            except Exception:
                pass

        try:
            page.locator("button[type='submit']").click(timeout=8000)
        except Exception:
            try:
                page.locator(
                    "button:has-text('Iniciar'), button:has-text('Login'), button:has-text('Ingresar')"
                ).first.click(timeout=5000)
            except Exception:
                try:
                    page.locator("button").first.click(timeout=8000)
                except Exception:
                    page.keyboard.press("Enter")

        page.wait_for_timeout(6000)

        if "/login" not in page.url:
            return

        if intento < max_intentos:
            page.wait_for_timeout(5000)
        else:
            raise Exception(f"Login C21 fallido tras {max_intentos} intentos — URL: {page.url}")


def exportar_reporte(page, fecha_desde: str, fecha_hasta: str) -> Path:
    """
    Usa el botón nativo "Descargar" (verde, debajo de los filtros) —
    confirmado que exporta el Excel completo con todas las columnas
    (PRECIO CIERRE, M2T, M2C, LAT, LON, etc.), sin necesitar ningún script
    de consola.

    Selector confirmado contra el sitio real: es un <a>, no un <button>
    (por eso get_by_role("button", ...) no lo encontraba).
    """
    url = construir_url_busqueda(fecha_desde, fecha_hasta)
    # La página puede tardar bastante (confirmado: hasta ~110 segundos,
    # con el ícono de carga girando en la pestaña) y a veces hace una
    # recarga automática antes de mostrar los resultados finales.
    # Timeout generoso a propósito (150s), con margen extra sobre lo medido.
    page.goto(url, wait_until="networkidle", timeout=150000)
    page.wait_for_selector("#myTableBusqueda tbody tr", timeout=150000)

    with page.expect_download(timeout=60000) as descarga_info:
        page.locator("a.btnDescargarInventarioExcel").click()
    descarga = descarga_info.value
    descarga.save_as(DESCARGA_TMP)
    return DESCARGA_TMP


def parse_money_usd(s):
    if not s:
        return None
    s = str(s)
    es_pyg = "PYG" in s.upper()
    s = s.replace("USD", "").replace("PYG", "").replace("Gs.", "").replace("Gs", "").strip()
    s = s.replace(".", "").replace(",", ".")
    try:
        val = float(s)
    except ValueError:
        return None
    if es_pyg:
        val = val / TC_APROX
    return val


def to_float(v):
    try:
        return float(v)
    except (TypeError, ValueError):
        return None


def validar_excel(path: Path) -> tuple[bool, str]:
    """Chequea que el archivo descargado sea un Excel legible con las columnas
    esperadas, antes de parsearlo. Devuelve (ok, motivo)."""
    if not path.exists() or path.stat().st_size == 0:
        return False, "El archivo descargado esta vacio o no se genero."
    try:
        wb = openpyxl.load_workbook(path, data_only=True)
        ws = wb.active
        header = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), None)
    except Exception as e:
        return False, f"El archivo descargado esta corrupto o no es un Excel valido: {e}"
    if not header:
        return False, "El archivo descargado no tiene fila de encabezado."
    faltantes = COLUMNAS_REQUERIDAS - {h for h in header if h}
    if faltantes:
        return False, f"Faltan columnas esperadas en el Excel descargado: {sorted(faltantes)}"
    return True, ""


def parsear_excel(path: Path) -> list[dict]:
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb.active  # el nombre de la hoja varia segun la exportacion (ej. "Hoja1")
    rows = list(ws.iter_rows(values_only=True))
    header = rows[0]
    idx = {h: i for i, h in enumerate(header) if h}

    out = []
    for r in rows[1:]:
        precio = parse_money_usd(r[idx["PRECIO CIERRE"]])
        precio_lista = parse_money_usd(r[idx["PRECIO"]])
        m2t = to_float(r[idx["M2T"]])
        m2c = to_float(r[idx["M2C"]])
        fc = r[idx["FECHA CIERRE"]]
        fa = r[idx["FECHA ALTA"]]
        pid = r[idx["ID"]]
        lat = to_float(r[idx["LAT"]])
        lon = to_float(r[idx["LON"]])
        if precio is None or not m2t or m2t <= 0 or not fc or not pid:
            continue
        out.append({
            "id": int(pid),
            "mu": r[idx["MUNICIPIO"]],
            "ti": r[idx["TIPO"]],
            "su": r[idx["SUBTIPO"]],
            "dir": r[idx["DIRECCION"]],
            "pr": round(precio, 2),
            "prL": round(precio_lista, 2) if precio_lista else None,
            "mt": m2t,
            "mc": m2c,
            "fc": str(fc)[:10],
            "fa": str(fa)[:10] if fa else None,
            "la": lat,
            "lo": lon,
        })
    return out


def refrescar_anho_actual_bcp():
    """Actualiza el año en curso de la tabla del BCP antes de usarla."""
    tabla = json.loads(TABLA_PATH.read_text(encoding="utf-8")) if TABLA_PATH.exists() else {}
    anho = datetime.now().year
    try:
        tabla.update(descargar_anho(anho))
        tabla = dict(sorted(tabla.items()))
        TABLA_PATH.write_text(json.dumps(tabla, ensure_ascii=False, indent=1), encoding="utf-8")
        print(f"Tabla BCP refrescada para {anho}: {len(tabla)} días")
    except Exception as e:
        print(f"AVISO: no se pudo refrescar la tabla BCP ({e}). Se usa la que hay.")
    return tabla


def normalizar_precios(nuevos, dataset_actual, tabla_bcp):
    """
    Aplica la misma corrección de moneda a los registros entrantes.
    Devuelve (nuevos_normalizados, conteos, diagnosticos).

    `upsert` reemplaza el registro entero por id — no hace merge de campos.
    Por eso la rama `marcar` es especial: si `clasificar` no pudo derivar
    nada confiable (el entrante trae el `pr` crudo en Gs, conocidamente
    basura) y el id ya existe en `dataset_actual` con un `pr` plausible
    (ya en USD, `< UMBRAL_GS`), se preserva ese `pr` almacenado en vez del
    crudo — y si nadie le sacó el `px` a mano todavía, se lo deja puesto;
    si ya no tiene `px`, es una resolución manual y no se la vuelve a
    marcar. Esto evita que un registro ambiguo, revisado y corregido a
    mano, sea pisado por su propia versión cruda si reaparece dentro de
    la ventana de 14 días del export semanal. Las demás ramas
    (`sin_cambio`, `corregir`, `corregir_y_marcar`, `eliminar`) no
    cambian: ahí el valor entrante ya es confiable o se sabe derivar, y
    `upsert` tiene que seguir propagando los cambios reales de la fuente.
    """
    pool = construir_pool(dataset_actual)
    por_id_actual = {d["id"]: d for d in dataset_actual}
    salida, diagnosticos = [], []
    conteos = {"corregidos": 0, "eliminados": 0, "marcados": 0,
               "preservados": 0, "errores": 0}
    for reg in nuevos:
        try:
            r = clasificar(reg, tabla_bcp, pool)
        except Exception as e:
            # Un fc malformado (u otro dato inesperado) no puede tumbar la
            # corrida semanal entera. Se deja pasar sin tocar el pr y se
            # marca para revisión manual, contado aparte.
            reg["px"] = 1
            conteos["errores"] += 1
            print(f"  ERROR clasificando id={reg['id']}: {type(e).__name__}: {e}")
            diagnosticos.append({"id": reg["id"], "accion": "error",
                                  "causa": type(e).__name__, "detalle": str(e)})
            salida.append(reg)
            continue

        accion = r["accion"]

        if accion == "eliminar":
            conteos["eliminados"] += 1
            print(f"  ELIMINADO id={reg['id']}: {r['detalle']}")
            diagnosticos.append({"id": reg["id"], "accion": accion,
                                  "causa": r["causa"], "detalle": r["detalle"]})
            continue

        if accion in ("corregir", "corregir_y_marcar"):
            reg["pr"] = r["pr_nuevo"]
            conteos["corregidos"] += 1

        if accion == "corregir_y_marcar":
            reg["px"] = 1
            conteos["marcados"] += 1
            print(f"  MARCADO id={reg['id']}: {r['causa']} — {r['detalle']}")
            diag_accion, diag_causa, diag_detalle = accion, r["causa"], r["detalle"]
        elif accion == "marcar":
            almacenado = por_id_actual.get(reg["id"])
            preservar = (almacenado is not None
                         and almacenado.get("pr") is not None
                         and almacenado["pr"] < UMBRAL_GS)
            if preservar:
                reg["pr"] = almacenado["pr"]
                if almacenado.get("px"):
                    reg["px"] = 1
                # si almacenado NO tiene px, no se lo ponemos: es una
                # resolución manual y se respeta tal cual.
                conteos["preservados"] += 1
                print(f"  PRESERVADO id={reg['id']}: se conserva pr={almacenado['pr']} "
                      f"del dataset ({r['causa']})")
                diag_accion = "preservado"
                diag_causa = r["causa"]
                diag_detalle = f"pr almacenado preservado ({almacenado['pr']}); {r['detalle']}"
            else:
                reg["px"] = 1
                conteos["marcados"] += 1
                print(f"  MARCADO id={reg['id']}: {r['causa']} — {r['detalle']}")
                diag_accion, diag_causa, diag_detalle = accion, r["causa"], r["detalle"]
        else:
            # sin_cambio
            diag_accion, diag_causa, diag_detalle = accion, r["causa"], r["detalle"]

        diagnosticos.append({"id": reg["id"], "accion": diag_accion,
                              "causa": diag_causa, "detalle": diag_detalle})
        salida.append(reg)
    return salida, conteos, diagnosticos


def cargar_dataset_actual() -> list[dict]:
    if DATASET_PATH.exists():
        with open(DATASET_PATH, encoding="utf-8") as f:
            return json.load(f)
    return []


def fecha_de_corte(dataset: list[dict]) -> str:
    if not dataset:
        # sin dataset previo: primer corrida, usar un rango amplio por defecto
        return (datetime.now() - timedelta(days=548)).strftime("%Y-%m-%d")  # ~18 meses
    ultima = max(d["fc"] for d in dataset)
    corte = datetime.strptime(ultima, "%Y-%m-%d") - timedelta(days=MARGEN_DIAS)
    return corte.strftime("%Y-%m-%d")


def upsert(dataset: list[dict], nuevos: list[dict]) -> tuple[list[dict], int, int]:
    por_id = {d["id"]: d for d in dataset}
    agregados = 0
    actualizados = 0
    for n in nuevos:
        if n["id"] in por_id:
            if por_id[n["id"]] != n:
                actualizados += 1
            por_id[n["id"]] = n
        else:
            por_id[n["id"]] = n
            agregados += 1
    resultado = list(por_id.values())
    resultado.sort(key=lambda d: d["fc"])
    return resultado, agregados, actualizados


def git_commit_push(mensaje: str):
    subprocess.run(["git", "add", "data/cerradas.json"], cwd=REPO_DIR, check=True)
    resultado = subprocess.run(["git", "diff", "--cached", "--quiet"], cwd=REPO_DIR)
    if resultado.returncode == 0:
        print("Sin cambios para commitear.")
        return
    subprocess.run(["git", "commit", "-m", mensaje], cwd=REPO_DIR, check=True)
    subprocess.run(["git", "push"], cwd=REPO_DIR, check=True)


def escribir_resumen_excel(fecha: str, agregados: int, actualizados: int, total: int,
                            nuevos: list[dict], alerta: str | None,
                            conteos_precio: dict | None = None,
                            diagnosticos: list[dict] | None = None) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Resumen"
    ws.append(["Fecha de corrida", fecha])
    ws.append(["Propiedades nuevas", agregados])
    ws.append(["Propiedades actualizadas", actualizados])
    ws.append(["Total acumulado", total])
    if conteos_precio:
        ws.append(["Precios corregidos (Gs->USD)", conteos_precio["corregidos"]])
        ws.append(["Registros eliminados (no son venta)", conteos_precio["eliminados"]])
        ws.append(["Registros marcados para revisión", conteos_precio["marcados"]])
        ws.append(["Registros preservados (resolución manual respetada)",
                   conteos_precio.get("preservados", 0)])
        ws.append(["Registros con error de clasificación", conteos_precio.get("errores", 0)])
    ws.append(["Estado", f"ALERTA - push frenado: {alerta}" if alerta else "OK - commit y push realizados"])
    ws.append([])
    ws.append(["ID", "Direccion", "Barrio", "Precio", "Fecha de cierre"])
    for n in nuevos:
        ws.append([n["id"], n["dir"], n["mu"], n["pr"], n["fc"]])

    if diagnosticos:
        ws_diag = wb.create_sheet("Diagnóstico precios")
        ws_diag.append(["ID", "Acción", "Causa", "Detalle"])
        for d in diagnosticos:
            ws_diag.append([d["id"], d["accion"], d["causa"], d["detalle"]])

    wb.save(RESUMEN_PATH)


def main():
    dataset = cargar_dataset_actual()
    total_previo = len(dataset)
    fecha_desde = fecha_de_corte(dataset)
    fecha_hasta = datetime.now().strftime("%Y-%m-%d")
    fecha_hoy = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    print(f"Dataset actual: {total_previo} registros. Exportando desde {fecha_desde} hasta {fecha_hasta}.")

    with sync_playwright() as p:
        browser = p.chromium.launch(headless=True)
        page = browser.new_page()
        login(page)
        ruta_export = exportar_reporte(page, fecha_desde, fecha_hasta)
        browser.close()

    ok, motivo = validar_excel(ruta_export)
    if not ok:
        alerta = f"Export invalido, se aborto la corrida sin tocar el dataset: {motivo}"
        print(f"ALERTA: {alerta}")
        escribir_resumen_excel(fecha_hoy, 0, 0, total_previo, [], alerta=alerta)
        ruta_export.unlink(missing_ok=True)
        return

    nuevos = parsear_excel(ruta_export)
    print(f"Exportados {len(nuevos)} registros en el rango.")

    tabla_bcp = refrescar_anho_actual_bcp()
    nuevos, conteos_precio, diagnosticos_precio = normalizar_precios(nuevos, dataset, tabla_bcp)
    print(f"Normalización de precios: {conteos_precio['corregidos']} corregidos, "
          f"{conteos_precio['eliminados']} eliminados, {conteos_precio['marcados']} marcados, "
          f"{conteos_precio['preservados']} preservados, {conteos_precio['errores']} con error.")

    dataset_final, agregados, actualizados = upsert(dataset, nuevos)
    total_final = len(dataset_final)

    alerta = None
    if agregados > LIMITE_NUEVAS:
        alerta = f"Las propiedades nuevas ({agregados}) superan el limite de seguridad ({LIMITE_NUEVAS})."
    elif total_final < total_previo:
        alerta = f"El total de registros bajo de {total_previo} a {total_final} respecto a la corrida anterior."

    DATASET_PATH.parent.mkdir(parents=True, exist_ok=True)
    with open(DATASET_PATH, "w", encoding="utf-8") as f:
        json.dump(dataset_final, f, ensure_ascii=False, separators=(",", ":"))

    print(f"Agregados: {agregados} | Actualizados: {actualizados} | Total dataset: {total_final}")

    if alerta:
        print(f"ALERTA: {alerta} No se hizo push automatico — revisar manualmente antes de pushear.")
    else:
        git_commit_push(f"Actualización semanal cerradas {fecha_hoy[:10]}: +{agregados} nuevos, {actualizados} actualizados")

    escribir_resumen_excel(fecha_hoy, agregados, actualizados, total_final, nuevos,
                           alerta=alerta, conteos_precio=conteos_precio,
                           diagnosticos=diagnosticos_precio)
    ruta_export.unlink(missing_ok=True)


if __name__ == "__main__":
    main()