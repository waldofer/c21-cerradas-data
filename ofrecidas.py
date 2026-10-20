"""Propiedades ofrecidas (a la venta) exportadas de 21online.

Modulo puro: solo openpyxl y la libreria estandar. No importa
`actualizar_cerradas` (ese arrastra playwright).

Privacidad: el export trae datos del propietario y columnas comerciales. Nada de
eso puede salir de aca; la salida se arma SIEMPRE desde la lista blanca `CAMPOS`.
"""
import json
import math
import re
import sys
import unicodedata
from pathlib import Path

import openpyxl

# Orden y lista blanca de claves de salida.
CAMPOS = ("id", "mu", "ti", "su", "dir", "en", "of", "pp", "mon", "mt", "mc", "fa", "la", "lo")

COLUMNAS_REQUERIDAS = {"ID", "PRECIO", "M2T", "FECHA ALTA", "MUNICIPIO"}

# Columna del export -> clave de salida (texto sin conversion).
_TEXTOS = {
    "mu": "MUNICIPIO",
    "ti": "TIPO",
    "su": "SUBTIPO",
    "dir": "DIRECCION",
    "en": "ENCABEZADO",
    "of": "OFICINA",
}

# Moneda: "USD", "PYG" o "Gs"/"Gs." (guaranies), antes o despues del numero.
_RE_MONEDA = re.compile(r"\b(USD|PYG|GS)\b\.?", re.IGNORECASE)


def parse_precio(s):
    """'420.000.000,00 PYG' -> (420000000.0, 'PYG'). None si no hay numero o moneda."""
    if s is None:
        return None
    s = str(s).strip()
    m = _RE_MONEDA.search(s)
    if not m:
        return None
    moneda = "USD" if m.group(1).upper() == "USD" else "PYG"
    numero = (s[:m.start()] + " " + s[m.end():]).strip()
    # Formato paraguayo: punto de miles, coma decimal.
    numero = numero.replace(".", "").replace(",", ".")
    try:
        valor = float(numero)
    except ValueError:
        return None
    return (valor, moneda) if math.isfinite(valor) else None  # sin NaN / inf


def validar_columnas(header):
    """Motivo (str) si falta alguna columna requerida; None si esta todo."""
    presentes = {str(h).strip() for h in header if h}
    faltantes = COLUMNAS_REQUERIDAS - presentes
    if faltantes:
        return f"Faltan columnas esperadas en el Excel de ofrecidas: {sorted(faltantes)}"
    return None


# Separador entre grupos de digitos.
_S = r"[\s.\-]"
# Cuerpo de un telefono: 6 digitos en 3+3 o en 2+2+2, que no sean cola de otro numero.
_CUERPO = rf"(?:{_S}?\d{{3}}{_S}?\d{{3}}|{_S}?\d{{2}}{_S}?\d{{2}}{_S}?\d{{2}})(?!\d)"
# Igual pero sin puntos (para formas sin prefijo, que se confunden con precios "Gs. 1.981.123.456").
_CUERPO_SIN_PUNTOS = r"(?:[\s-]?\d{3}[\s-]?\d{3}|[\s-]?\d{2}[\s-]?\d{2}[\s-]?\d{2})(?!\d)"
# Ni pegado a otro numero ni como tramo de un numero agrupado ("1 050 000 000", "1.025.000.000").
_SUELTO = r"(?<!\d)(?<!\d[\s.,\-])"

# Patrones que se borran del texto libre, en orden. Cada uno arranca comiendo los separadores
# previos para no dejar guiones colgando ("Casa 12-0981-123-456" -> "Casa 12").
_PATRONES_TELEFONO = tuple(re.compile(p) for p in (
    # Emails.
    r"[\s-]*\S+@\S+\.\w+",
    # Celulares con prefijo (+595 / 595 / 0, con o sin parentesis): basta que no sea cola de otro
    # numero pegado ("Mz 5 0981 123 456" -> "Mz 5").
    rf"[\s-]*(?<!\d)\(?(?:\+?595{_S}?|0)9\d{{2}}\)?{_CUERPO}",
    # Celulares sin prefijo: solo 96x-99x, sueltos y sin puntos.
    rf"[\s-]*{_SUELTO}9[6-9]\d{_CUERPO_SIN_PUNTOS}",
    # Fijos con codigo de area y 0 (021, 061, 0521...), sueltos.
    rf"[\s-]*{_SUELTO}\(?0[2-8]\d{{1,2}}\)?{_CUERPO}",
    # Fijos con +595 / 595 y codigo de area sin 0 (21, 61, 521...).
    rf"[\s-]*(?<![\d.,])\(?\+?595\)?{_S}?\(?0?[2-8]\d{{1,2}}\)?{_CUERPO}",
    # Fijo de Asuncion sin prefijo: 21 + 6 digitos, suelto y sin puntos ni comas que lo
    # pegen a otro numero ("Gs.21612345", "21612345,50" quedan; un punto final no molesta).
    rf"[\s-]*{_SUELTO}(?<![.,])21[\s-]?\d{{3}}[\s-]?\d{{3}}(?!\d)(?![.,]\d)",
    # Cualquier corrida de 9 o mas digitos seguidos.
    r"[\s-]*\d{9,}",
))


def quitar_telefonos(v):
    """Saca telefonos del texto libre (DIRECCION, ENCABEZADO). None si el campo queda vacio.

    Saca celulares, fijos y emails. No toca precios con puntos de miles ni numeros cortos.
    """
    if v is None:
        return None
    for patron in _PATRONES_TELEFONO:
        v = patron.sub(" ", v)
    v = re.sub(r"\s+", " ", v).strip()
    return v or None


# --- Barrera antes de publicar ---
# El filtro de arriba borra lo que reconoce. Esta barrera es mas desconfiada: si en dir/en queda
# algo con pinta de telefono (aunque sea una forma que el filtro no conoce), no se publica nada.
# Preferimos frenar una semana a tener que reescribir el historial del repo publico.

# Todo lo que no sea un digito ni una letra y tenga hasta 3 caracteres se toma como separador
# entre grupos de un mismo numero ("0981 123 456", "0981/123456", "1234,021 612 345",
# "0981–123–456"). Ante la duda, se juntan: la barrera prefiere frenar de mas.
_RE_GRUPO = re.compile(r"\d+")
_RE_SEPARADOR = re.compile(r"[^\w@]{1,3}")
# Codigo postal paraguayo (6 digitos que empiezan con 0), que el geocodificador agrega a la
# direccion: "Asuncion 001204 388", "SAN JUAN BAUTISTA 080106 080106".
_RE_CODIGO_POSTAL = re.compile(r"0\d{5}")
# Un telefono tiene a lo sumo 13 digitos (+595 981 123 456): las ventanas mas largas no cuentan.
_MAX_GRUPOS_VENTANA = 6


def _secuencias(texto: str):
    """Grupos de digitos consecutivos separados solo por separadores cortos.

    Devuelve listas de (grupo, separador_previo)."""
    secuencia, fin_anterior = [], None
    for m in _RE_GRUPO.finditer(texto):
        sep = texto[fin_anterior:m.start()] if fin_anterior is not None else None
        if secuencia and not _RE_SEPARADOR.fullmatch(sep):
            yield secuencia
            secuencia = []
        secuencia.append((m.group(), sep if secuencia else None))
        fin_anterior = m.end()
    if secuencia:
        yield secuencia


def _es_monto(secuencia) -> bool:
    """Monto redondo con miles: "1 050 000 000", "420.000.000". Los precios terminan en 000."""
    grupos = [g for g, _ in secuencia]
    seps = {s.strip() or " " for _, s in secuencia[1:]}
    return (
        len(grupos) >= 2 and len(grupos[0]) <= 3 and all(len(g) == 3 for g in grupos[1:])
        and len(seps) == 1 and seps <= {".", " "} and grupos[-1] == "000"
    )


def _ventana_parece_telefono(ventana) -> bool:
    grupos = [g for g, _ in ventana]
    if any(_RE_CODIGO_POSTAL.fullmatch(g) for g in grupos):
        return False
    # Coordenada decimal: "-25.2674758", "-59.5123456".
    if len(grupos) >= 2 and len(grupos[0]) <= 2 and ventana[1][1] == "." and len(grupos[1]) >= 5:
        return False
    d = "".join(grupos)
    return (
        (9 <= len(d) <= 13 and (d.startswith("0") or d.startswith("595")))  # con prefijo
        or (len(d) == 9 and d[0] == "9" and d[1] in "6789")  # celular sin prefijo
        or (len(d) == 8 and d.startswith("21"))  # fijo de Asuncion sin el 0
    )


def parece_telefono(texto) -> bool:
    """True si el texto tiene un email o algo que podria ser un telefono."""
    if not texto:
        return False
    # Digitos de ancho completo ("０９８１") pasan a digitos comunes.
    texto = unicodedata.normalize("NFKC", texto)
    if "@" in texto or re.search(r"\d{9,}", texto):
        return True
    for secuencia in _secuencias(texto):
        if _es_monto(secuencia):
            continue
        # El telefono puede empezar en cualquier grupo ("Lote 3 981 123 456 USD 120.000").
        for i in range(len(secuencia)):
            for j in range(i + 1, min(len(secuencia), i + _MAX_GRUPOS_VENTANA) + 1):
                if _ventana_parece_telefono(secuencia[i:j]):
                    return True
    return False


def telefonos_residuales(registros: list[dict]) -> list[int]:
    """IDs de los registros cuya direccion o titulo tiene algo con pinta de telefono."""
    return [r["id"] for r in registros if parece_telefono(r.get("dir")) or parece_telefono(r.get("en"))]


def _lista_ids(ids: list[int], tope: int = 20) -> str:
    return ", ".join(str(i) for i in ids[:tope]) + (" ..." if len(ids) > tope else "")


def _texto(v):
    """Celda -> str sin espacios, o None si esta vacia."""
    if v is None:
        return None
    v = str(v).strip()
    return v or None


def _numero(v):
    """Celda de texto/numero -> float, o None si esta vacia, no es un numero o no es finito."""
    try:
        valor = float(_texto(v))
    except (TypeError, ValueError):
        return None
    return valor if math.isfinite(valor) else None


def parsear_ofrecidas(path: Path) -> list[dict]:
    """Un dict por ID (gana la ultima fila), solo con las claves de CAMPOS.

    Se descartan las filas sin ID, con precio ilegible o con M2T faltante o <= 0.
    """
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb.active  # el nombre de la hoja varia segun la exportacion
        filas = ws.iter_rows(values_only=True)
        header = next(filas, None) or ()
        motivo = validar_columnas(header)
        if motivo:
            raise ValueError(motivo)
        idx = {str(h).strip(): i for i, h in enumerate(header) if h}

        def celda(fila, col):
            i = idx.get(col)
            return fila[i] if i is not None and i < len(fila) else None

        por_id = {}
        for fila in filas:
            pid = _numero(celda(fila, "ID"))
            precio = parse_precio(celda(fila, "PRECIO"))
            m2t = _numero(celda(fila, "M2T"))
            if pid is None or precio is None or m2t is None or m2t <= 0:
                continue
            fa = _texto(celda(fila, "FECHA ALTA"))
            valores = {
                "id": int(pid),
                "pp": precio[0],
                "mon": precio[1],
                "mt": m2t,
                "mc": _numero(celda(fila, "M2C")),
                "fa": fa[:10] if fa else None,
                "la": _numero(celda(fila, "LAT")),
                "lo": _numero(celda(fila, "LON")),
            }
            for clave, col in _TEXTOS.items():
                valores[clave] = _texto(celda(fila, col))
            # Texto libre: el agente a veces escribe telefonos adentro.
            valores["dir"] = quitar_telefonos(valores["dir"])
            valores["en"] = quitar_telefonos(valores["en"])
            # Lista blanca: solo salen las claves de CAMPOS.
            por_id[valores["id"]] = {k: valores[k] for k in CAMPOS}
        return list(por_id.values())
    finally:
        wb.close()


# Si la descarga trae menos del 80% de las ofrecidas de la semana anterior,
# se la trata como rota y no se publica.
UMBRAL_DESCARGA = 0.8


def descarga_sospechosa(n_nuevas: int, n_anteriores: int, umbral: float = UMBRAL_DESCARGA):
    """Mensaje (str) si la descarga parece cortada; None si es normal o es la primera corrida."""
    if n_anteriores == 0 or n_nuevas >= umbral * n_anteriores:
        return None
    return (
        f"La descarga de ofrecidas trajo {n_nuevas} propiedades y la semana anterior "
        f"habia {n_anteriores} (menos del {umbral:.0%}); se la considera incompleta."
    )


def comparar(anteriores: list[dict], descargadas: list[dict], vendidas_por_id: dict[int, dict]) -> dict:
    """Compara la descarga de esta semana con la anterior y con las vendidas.

    Una ofrecida que ya figura en `vendidas_por_id` no se publica como ofrecida.
    """
    ant = {r["id"]: r for r in anteriores}
    desc = {r["id"]: r for r in descargadas}

    ofrecidas = sorted(
        (r for i, r in desc.items() if i not in vendidas_por_id), key=lambda r: r["id"]
    )
    nuevas = [r for r in ofrecidas if r["id"] not in ant]
    pasaron_a_vendidas = [
        {
            "id": i,
            "dir": r["dir"],
            "mu": r["mu"],
            "pp": r["pp"],
            "mon": r["mon"],
            "pr_cierre": vendidas_por_id[i]["pr"],
        }
        for i, r in ant.items()
        if i in vendidas_por_id
    ]
    dadas_de_baja = [
        r for i, r in ant.items() if i not in desc and i not in vendidas_por_id
    ]
    cambios_precio = [
        {
            "id": i,
            "antes": (ant[i]["pp"], ant[i]["mon"]),
            "despues": (r["pp"], r["mon"]),
        }
        for i, r in desc.items()
        if i in ant and (ant[i]["pp"], ant[i]["mon"]) != (r["pp"], r["mon"])
    ]
    return {
        "ofrecidas": ofrecidas,
        "nuevas": nuevas,
        "pasaron_a_vendidas": pasaron_a_vendidas,
        "dadas_de_baja": dadas_de_baja,
        "cambios_precio": cambios_precio,
        "en_venta_antes": len(anteriores),
    }


OFRECIDAS_PATH = Path(__file__).parent / "data" / "ofrecidas.json"
CERRADAS_PATH = Path(__file__).parent / "data" / "cerradas.json"


def guardar_ofrecidas(registros: list[dict], path: Path = OFRECIDAS_PATH) -> None:
    """JSON compacto ordenado por id. Escribe a .tmp y reemplaza, para no dejar un archivo roto."""
    path = Path(path)
    # Ultimo cerrojo: aunque alguien llame a esto sin pasar por procesar_ofrecidas (ej. `inicial`).
    sospechosos = telefonos_residuales(registros)
    if sospechosos:
        raise ValueError(
            f"No se guarda ofrecidas.json: {len(sospechosos)} propiedades tienen algo con pinta de "
            f"telefono en la direccion o el titulo. IDs: {_lista_ids(sospechosos)}"
        )
    ordenados = sorted(registros, key=lambda r: r["id"])
    tmp = path.with_name(path.name + ".tmp")
    # allow_nan=False: un NaN/inf no es JSON valido; mejor fallar que publicar un archivo roto.
    texto = json.dumps(ordenados, ensure_ascii=False, separators=(",", ":"), allow_nan=False)
    tmp.write_text(texto, encoding="utf-8")
    tmp.replace(path)


def _cerradas_por_id(path: Path = CERRADAS_PATH) -> dict[int, dict]:
    with open(path, encoding="utf-8") as f:
        return {r["id"]: r for r in json.load(f)}


def _filas_de_datos(ruta: Path) -> int:
    """Cantidad de filas con algun valor debajo del encabezado."""
    wb = openpyxl.load_workbook(ruta, read_only=True, data_only=True)
    try:
        filas = wb.active.iter_rows(values_only=True)
        next(filas, None)
        return sum(1 for f in filas if any(c is not None and str(c).strip() for c in f))
    finally:
        wb.close()


def _header(ruta: Path) -> tuple:
    """Fila de encabezado del Excel (tupla vacia si no hay)."""
    wb = openpyxl.load_workbook(ruta, read_only=True, data_only=True)
    try:
        return next(wb.active.iter_rows(values_only=True), None) or ()
    finally:
        wb.close()


# Estatus que tiene que traer cada fila del export de ofrecidas.
ESTATUS_ESPERADO = "enPromocion"
# Mas nuevas que esto en una semana (con semana anterior) es una descarga inflada.
LIMITE_NUEVAS_OFRECIDAS = 1500


def estatus_inesperados(ruta: Path) -> int:
    """Cantidad de filas con ID cuyo ESTATUS no es enPromocion (o falta la columna).

    El ESTATUS solo se usa para validar; no sale del modulo.
    """
    wb = openpyxl.load_workbook(ruta, read_only=True, data_only=True)
    try:
        filas = wb.active.iter_rows(values_only=True)
        header = next(filas, None) or ()
        idx = {str(h).strip(): i for i, h in enumerate(header) if h}
        i_id, i_est = idx.get("ID"), idx.get("ESTATUS")
        n = 0
        for fila in filas:
            if i_id is None or i_id >= len(fila) or _texto(fila[i_id]) is None:
                continue
            estatus = _texto(fila[i_est]) if i_est is not None and i_est < len(fila) else None
            if estatus != ESTATUS_ESPERADO:
                n += 1
        return n
    finally:
        wb.close()


def procesar_ofrecidas(
    ruta_export: Path, anteriores: list[dict], vendidas_por_id: dict[int, dict]
) -> tuple[list[dict] | None, dict | None, str | None]:
    """Parsea el export y lo cruza con la semana anterior y con las vendidas.

    Devuelve (final, resultado, alerta). Si hay alerta, final es None y no se debe escribir nada.
    """
    motivo = validar_columnas(_header(ruta_export))
    if motivo:
        return None, None, motivo
    try:
        descargadas = parsear_ofrecidas(ruta_export)
    except ValueError as e:
        return None, None, str(e)
    if not descargadas:
        return None, None, "El export de ofrecidas no trajo ninguna propiedad valida."
    otros = estatus_inesperados(ruta_export)
    if otros:
        return None, None, (
            f"El export de ofrecidas trae {otros} filas con estatus distinto de "
            f"{ESTATUS_ESPERADO}; no se publica."
        )
    alerta = descarga_sospechosa(len(descargadas), len(anteriores))
    if alerta:
        return None, None, alerta
    resultado = comparar(anteriores, descargadas, vendidas_por_id)
    n_nuevas = len(resultado["nuevas"])
    if anteriores and n_nuevas > LIMITE_NUEVAS_OFRECIDAS:
        return None, resultado, (
            f"Las ofrecidas nuevas ({n_nuevas}) superan el limite de seguridad "
            f"({LIMITE_NUEVAS_OFRECIDAS}); se la considera una descarga inflada."
        )
    sospechosos = telefonos_residuales(resultado["ofrecidas"])
    if sospechosos:
        return None, resultado, (
            f"{len(sospechosos)} propiedades tienen en la direccion o el titulo algo con pinta de "
            f"telefono que el filtro no limpio; esta semana no se publica. Revisar estos IDs en "
            f"21online y ampliar el filtro: {_lista_ids(sospechosos)}"
        )
    return resultado["ofrecidas"], resultado, None


def escribir_hoja_ofrecidas(ruta_resumen: Path, resultado: dict | None, alerta: str | None) -> None:
    """Crea o reemplaza la hoja "Ofrecidas" de resumen_actualizacion.xlsx (crea el libro si falta)."""
    ruta_resumen = Path(ruta_resumen)
    if ruta_resumen.exists():
        wb = openpyxl.load_workbook(ruta_resumen)
        if "Ofrecidas" in wb.sheetnames:
            del wb["Ofrecidas"]
        ws = wb.create_sheet("Ofrecidas")
    else:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Ofrecidas"

    ws.append(["Estado", f"ALERTA - {alerta}" if alerta else "OK"])
    if resultado:
        ws.append(["En venta antes", resultado["en_venta_antes"]])
        ws.append(["Nuevas", len(resultado["nuevas"])])
        ws.append(["Pasaron a vendidas", len(resultado["pasaron_a_vendidas"])])
        ws.append(["Dadas de baja", len(resultado["dadas_de_baja"])])
        ws.append(["En venta ahora", len(resultado["ofrecidas"])])
        ws.append([])
        ws.append(["Pasaron a vendidas"])
        ws.append(["ID", "Dirección", "Barrio", "Precio pedido", "Moneda", "Precio de cierre USD"])
        vendidas = set()
        for r in resultado["pasaron_a_vendidas"]:
            vendidas.add(r["id"])
            ws.append([r["id"], r["dir"], r["mu"], r["pp"], r["mon"], r["pr_cierre"]])
        ws.append([])
        ws.append(["Cambios de precio"])
        ws.append(["ID", "Antes", "Moneda", "Después", "Moneda"])
        for c in resultado["cambios_precio"]:
            if c["id"] in vendidas:  # ya figura como vendida: su cambio de precio es ruido
                continue
            ws.append([c["id"], c["antes"][0], c["antes"][1], c["despues"][0], c["despues"][1]])
    wb.save(ruta_resumen)


def _inicial(ruta: Path) -> None:
    leidas = _filas_de_datos(ruta)
    parseadas = parsear_ofrecidas(ruta)  # valida columnas; levanta ValueError si faltan
    publicables = comparar([], parseadas, _cerradas_por_id())["ofrecidas"]
    guardar_ofrecidas(publicables)
    print(
        f"Filas leidas: {leidas} | descartadas (sin M2T, precio ilegible, sin ID o repetidas): "
        f"{leidas - len(parseadas)} | ya vendidas (excluidas): {len(parseadas) - len(publicables)} "
        f"| guardadas: {len(publicables)}"
    )


if __name__ == "__main__":
    if len(sys.argv) == 3 and sys.argv[1] == "inicial":
        _inicial(Path(sys.argv[2]))
    else:
        sys.exit("Uso: python ofrecidas.py inicial <ruta.xlsx>")
